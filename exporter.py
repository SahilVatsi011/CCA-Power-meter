"""Export readings to CSV or Excel with resampling and field selection."""

import csv
import io
import math
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import openpyxl

import config
import db

IST = ZoneInfo(config.IST_ZONE)

FIELD_COLUMNS = ["time_ist", "ts_ms", "device", "v", "a", "w", "pf", "kwh", "temp", "leak"]
FIELD_LABELS = {
    "time": "Time (IST)",
    "device": "Device ID",
    "v": "Voltage (V)",
    "a": "Current (A)",
    "w": "Power (W)",
    "pf": "Power Factor",
    "kwh": "Energy (kWh)",
    "temp": "Temperature (°C)",
    "leak": "Leakage (mA)",
}

ALLOWED_FIELDS = ("v", "a", "w", "pf", "kwh", "temp", "leak")
WINDOW_SEC = {"raw": 0, "1min": 60, "5min": 300, "15min": 900, "hourly": 3600}
MAX_EXCEL_ROWS = 200000


def _pf(w, v, a):
    if w is None or v is None or a is None:
        return None
    if v <= 0 or a <= 0:
        return None
    return round(w / (v * a), 3)


def _to_row(ts, device_id, values):
    local = ts.astimezone(IST).replace(tzinfo=None)
    row = {
        "time_ist": local.strftime("%Y-%m-%d %H:%M:%S"),
        "ts_ms": int(ts.timestamp() * 1000),
        "device": device_id,
        "v": values.get("v"),
        "a": values.get("a"),
        "w": values.get("w"),
        "pf": values.get("pf"),
        "kwh": values.get("kwh"),
        "temp": values.get("temp"),
        "leak": values.get("leak"),
    }
    return row


def _doc_values(doc, step):
    if step == "raw":
        w = doc.get("w")
        return {
            "v": doc.get("v"),
            "a": doc.get("a"),
            "w": w,
            "pf": doc.get("pf", _pf(w, doc.get("v"), doc.get("a"))),
            "kwh": doc.get("kwh"),
            "temp": doc.get("temp"),
            "leak": doc.get("leak"),
        }
    w = doc.get("w_avg")
    return {
        "v": doc.get("v_avg"),
        "a": doc.get("a_avg"),
        "w": w,
        "pf": _pf(w, doc.get("v_avg"), doc.get("a_avg")),
        "kwh": doc.get("kwh"),
        "temp": doc.get("temp_avg"),
        "leak": doc.get("leak_max"),
    }


def _source_collection(interval):
    if interval == "raw":
        return config.SERVER_DB
    if interval in ("1min", "5min", "15min"):
        return config.SERVER_DB_1MIN
    return config.SERVER_DB_HOURLY


def _bucket(docs, window):
    buckets = {}
    order = []
    for doc in docs:
        key = int(doc["ts"].timestamp() // window) * window
        if key not in buckets:
            buckets[key] = {
                "ts": datetime.fromtimestamp(key, tz=timezone.utc), "n": 0,
                "v_sum": 0.0, "a_sum": 0.0, "w_sum": 0.0,
                "temp_sum": 0.0, "leak_max": 0.0, "kwh": None,
            }
            order.append(key)
        b = buckets[key]
        v, a, w = doc.get("v_avg"), doc.get("a_avg"), doc.get("w_avg")
        b["n"] += 1
        if v is not None:
            b["v_sum"] += v
        if a is not None:
            b["a_sum"] += a
        if w is not None:
            b["w_sum"] += w
        temp = doc.get("temp_avg")
        if temp is not None:
            b["temp_sum"] += temp
        le = doc.get("leak_max")
        if le is not None:
            b["leak_max"] = max(b["leak_max"], le)
        kwh = doc.get("kwh")
        if kwh is not None:
            b["kwh"] = kwh

    order.sort()
    rows = []
    for key in order:
        b = buckets[key]
        n = b["n"]
        rows.append(_to_row(
            b["ts"], b.get("device_id", ""),
            {
                "v": round(b["v_sum"] / n, 1) if b["v_sum"] else None,
                "a": round(b["a_sum"] / n, 3) if b["a_sum"] else None,
                "w": round(b["w_sum"] / n) if b["w_sum"] else None,
                "kwh": b["kwh"],
                "temp": round(b["temp_sum"] / n, 1) if b["temp_sum"] else None,
                "leak": b["leak_max"],
            },
        ))
    return rows


def iter_rows(from_ts, to_ts, interval):
    if interval not in WINDOW_SEC:
        raise ValueError(f"invalid interval: {interval}")

    coll = _source_collection(interval)
    docs = db.query_docs(coll, from_ts, to_ts)

    if interval in ("5min", "15min"):
        rows = _bucket(docs, WINDOW_SEC[interval])
        for row in rows:
            yield row
        return

    for doc in docs:
        step = doc.get("step", "raw" if coll == config.SERVER_DB else "minute")
        yield _to_row(doc["ts"], doc.get("device_id", ""), _doc_values(doc, step))


def _header_and_rows(rows, fields):
    ordered = ["time_ist"] + [f for f in FIELD_COLUMNS if f in fields and f != "time_ist"]
    header = [FIELD_LABELS.get(c, c) for c in ordered]
    out = []
    for row in rows:
        out.append([row.get(c) for c in ordered])
    return header, out


def build_csv(from_ts, to_ts, interval, fields):
    header, body = _header_and_rows(iter_rows(from_ts, to_ts, interval), fields)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(body)
    return buf.getvalue()


def build_xlsx(from_ts, to_ts, interval, fields):
    header, body = _header_and_rows(iter_rows(from_ts, to_ts, interval), fields)
    if len(body) > MAX_EXCEL_ROWS:
        body = body[:MAX_EXCEL_ROWS]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "readings"
    ws.append(header)
    for row in body:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


EVENT_LABELS = {
    "power_cut": "Power Cut",
    "power_restored": "Power Restored",
    "fault": "Fault",
    "fault_cleared": "Fault Cleared",
}


def event_rows(ev_list):
    rows = []
    for ev in ev_list:
        local = ev["ts"].astimezone(IST).replace(tzinfo=None)
        rows.append([
            local.strftime("%Y-%m-%d %H:%M:%S"),
            EVENT_LABELS.get(ev.get("event_type"), ev.get("event_type") or ""),
            ev.get("detail") or "",
        ])
    return rows


def build_events_csv(ev_list):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["time_ist", "event", "detail"])
    writer.writerows(event_rows(ev_list))
    return buf.getvalue()


def build_events_xlsx(ev_list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "events"
    ws.append(["time_ist", "event", "detail"])
    for row in event_rows(ev_list):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()